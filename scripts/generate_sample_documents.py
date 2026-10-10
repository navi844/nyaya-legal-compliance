from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parents[1]
DOCUMENTS_DIR = ROOT / 'documents'
DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)


def write_pdf(path: Path, title: str, sections: list[str]) -> None:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=letter, title=title)
    story = []
    story.append(Paragraph('Document ID: NYA-MSA-001', styles['BodyText']))
    story.append(Paragraph('Version: 1.0', styles['BodyText']))
    story.append(Paragraph('Effective Date: 2026-01-01', styles['BodyText']))
    story.append(Paragraph('Classification: FICTIONAL TRAINING SAMPLE', styles['BodyText']))
    story.append(Paragraph('This document is not legal advice and contains no real company data.', styles['BodyText']))
    story.append(Spacer(1, 12))
    for section in sections:
        story.append(Paragraph(section, styles['BodyText']))
        story.append(Spacer(1, 8))
    doc.build(story)


def write_docx(path: Path, doc_id: str, version: str, effective_date: str, paragraphs: list[str]) -> None:
    document = Document()
    document.add_paragraph(f'Document ID: {doc_id}')
    document.add_paragraph(f'Version: {version}')
    document.add_paragraph(f'Effective Date: {effective_date}')
    document.add_paragraph('Classification: FICTIONAL TRAINING SAMPLE')
    document.add_paragraph('This document is not legal advice and contains no real company data.')
    document.add_paragraph('The regulator, portal, and currency units are fictional.')
    for para in paragraphs:
        document.add_paragraph(para)
    document.save(path)


def write_html(path: Path, title: str, body_sections: list[str]) -> None:
    html_parts = [
        '<!DOCTYPE html>',
        '<html><head><meta charset="utf-8"><title>{}</title></head><body>'.format(title),
        '<h1>Document ID: NYA-LFAQ-003</h1>',
        '<p>Version: 2.0</p>',
        '<p>Effective Date: 2026-07-15</p>',
        '<p>Classification: FICTIONAL TRAINING SAMPLE</p>',
        '<p>This document is not legal advice and contains no real company data.</p>',
    ]
    for section in body_sections:
        html_parts.append(section)
    html_parts.append('</body></html>')
    path.write_text('\n'.join(html_parts), encoding='utf-8')


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=['audit_id', 'date', 'area_audited', 'finding', 'severity', 'status'])
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def write_json(path: Path, payload: dict) -> None:
    path.write_text(json.dumps(payload, indent=2), encoding='utf-8')


def generate_documents() -> None:
    pdf_sections = [
        'Section 8 — Limitation of Liability',
        'The aggregate liability cap is an amount equal to fees paid or payable under the applicable Order Form during the 12 months immediately preceding the event giving rise to the claim.',
        'Exclusions from the general cap must be described separately so retrieval does not confuse them with the main cap.',
        'Section 9 — Indemnification',
        'The Supplier will indemnify the customer for third-party intellectual property infringement claims arising from the services.',
        'The Supplier will indemnify the customer for bodily injury, death, or tangible property damage caused by negligence or wilful misconduct.',
        'The Supplier will indemnify the customer for breach of confidentiality or applicable data-protection obligations.',
        'Section 14 — Governing Law and Jurisdiction',
        'Governing law: laws of the fictional jurisdiction of Dakshin Pradesh.',
        'Exclusive jurisdiction: commercial courts of Navapura.',
    ]
    write_pdf(DOCUMENTS_DIR / 'NYAYA_MSA_Excerpt.pdf', 'NYAYA MSA Excerpt', pdf_sections)

    docx_paragraphs = [
        'Regulated providers must submit a quarterly compliance report.',
        'The report deadline is the 15th calendar day after the end of each calendar quarter.',
        'If the date falls on a public holiday, submission is due on the next business day.',
        'Late filing penalty is 25,000 fictional compliance units per calendar day.',
        'The penalty is capped at 500,000 fictional compliance units per report.',
        'Submissions go through the fictional DCCA Compliance Gateway.',
        'The Digital Communications Compliance Authority, abbreviated DCCA, is fictional.',
    ]
    write_docx(
        DOCUMENTS_DIR / 'NYAYA_Regulatory_Compliance_Circular.docx',
        'NYA-RCC-2026-04',
        '1.1',
        '2026-04-01',
        docx_paragraphs,
    )

    html_sections = [
        '<section><h2>Standard review</h2><p>Standard contract review: five business days.</p></section>',
        '<section><h2>Expedited review</h2><p>Expedited contract review: two business days.</p></section>',
        '<section><h2>Urgent review</h2><p>Urgent review: submit an Urgent Legal Review form and then contact the fictional duty counsel.</p><p>Fictional contact: Asha Verma.</p><p>Fictional email: asha.verma@example.invalid</p><p>Fictional phone: +91 00000 12345</p><p>Urgent requests must include the customer name, contract value band, required signature date, risk summary, and account owner.</p><p>Urgent classification does not guarantee same-day approval.</p></section>',
    ]
    write_html(DOCUMENTS_DIR / 'NYAYA_Legal_FAQ.html', 'Nyaya Legal FAQ', html_sections)

    csv_rows = [
        {'audit_id': 'AUD-2026-001', 'date': '2026-01-12', 'area_audited': 'Access provisioning', 'finding': 'User access recertification completed on time', 'severity': 'Low', 'status': 'Closed'},
        {'audit_id': 'AUD-2026-002', 'date': '2026-02-14', 'area_audited': 'Privileged access', 'finding': 'privileged-access review evidence incomplete', 'severity': 'High', 'status': 'Open'},
        {'audit_id': 'AUD-2026-003', 'date': '2026-03-04', 'area_audited': 'Third-party onboarding', 'finding': 'Vendor risk questionnaire reviewed and approved', 'severity': 'Medium', 'status': 'Closed'},
        {'audit_id': 'AUD-2026-004', 'date': '2026-03-29', 'area_audited': 'Vendor due diligence', 'finding': 'vendor due-diligence renewal overdue', 'severity': 'Medium', 'status': 'Open'},
        {'audit_id': 'AUD-2026-005', 'date': '2026-04-18', 'area_audited': 'Data retention', 'finding': 'Retention schedule updated after remediation', 'severity': 'Low', 'status': 'Closed'},
        {'audit_id': 'AUD-2026-006', 'date': '2026-05-12', 'area_audited': 'Retention exception', 'finding': 'retention exception lacks documented approval', 'severity': 'High', 'status': 'Open'},
        {'audit_id': 'AUD-2026-007', 'date': '2026-06-08', 'area_audited': 'Incident response', 'finding': 'Incident drill report accepted by operating committee', 'severity': 'Medium', 'status': 'Closed'},
        {'audit_id': 'AUD-2026-008', 'date': '2026-07-19', 'area_audited': 'Business continuity', 'finding': 'incident exercise action owner not assigned', 'severity': 'Low', 'status': 'Open'},
    ]
    write_csv(DOCUMENTS_DIR / 'NYAYA_Compliance_Audit_Log.csv', csv_rows)

    filing_payload = {
        'document_id': 'NYA-RFT-010',
        'version': 1.0,
        'effective_date': '2026-10-10',
        'filings': [
            {'filing_id': 'FIL-2026-101', 'jurisdiction': 'Dakshin Pradesh', 'description': 'Annual network resilience attestation', 'due_date': '2026-10-05', 'status': 'Overdue'},
            {'filing_id': 'FIL-2026-102', 'jurisdiction': 'Navapura', 'description': 'Consumer complaint metrics return', 'due_date': '2026-10-25', 'status': 'Due'},
            {'filing_id': 'FIL-2026-103', 'jurisdiction': 'Paschim Territory', 'description': 'Spectrum usage declaration', 'due_date': '2026-11-05', 'status': 'Planned'},
            {'filing_id': 'FIL-2026-104', 'jurisdiction': 'Uttara Region', 'description': 'Emergency communications certification', 'due_date': '2026-09-30', 'status': 'Submitted'},
            {'filing_id': 'FIL-2026-105', 'jurisdiction': 'Dakshin Pradesh', 'description': 'Quarterly privacy controls statement', 'due_date': '2026-10-15', 'status': 'Due'},
            {'filing_id': 'FIL-2026-106', 'jurisdiction': 'Navapura', 'description': 'Lawful request transparency schedule', 'due_date': '2026-12-15', 'status': 'Planned'},
        ],
    }
    write_json(DOCUMENTS_DIR / 'NYAYA_Regulatory_Filing_Tracker.json', filing_payload)


if __name__ == '__main__':
    generate_documents()
    print('Generated sample documents at', DOCUMENTS_DIR)
